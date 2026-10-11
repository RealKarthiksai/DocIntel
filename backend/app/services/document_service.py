import os
import io
import uuid
import datetime
import logging
from typing import List, Dict, Any, Tuple
import pymupdf
import docx

from app.core.config import settings
from app.core.database import get_mongo_db, get_chroma_collection, save_mongo_mock_backup
from app.models.postgres_models import Document, DocumentChunk, DocumentAccessPolicy
from app.services.audit_service import log_audit_event
from sqlalchemy import func

logger = logging.getLogger(__name__)

_embedding_model = None

def get_embedding_model():
    global _embedding_model
    if _embedding_model is None:
        from sentence_transformers import SentenceTransformer
        _embedding_model = SentenceTransformer(settings.EMBEDDING_MODEL)
    return _embedding_model

def extract_text_from_file(filename: str, content: bytes) -> List[Tuple[int, str]]:
    """Extracts text page-by-page from PDF, DOCX, TXT, or Markdown. Returns list of (page_num, text)."""
    ext = os.path.splitext(filename)[1].lower()
    extracted_pages: List[Tuple[int, str]] = []
    
    if ext == ".pdf":
        doc = pymupdf.open(stream=content, filetype="pdf")
        for i in range(len(doc)):
            text = doc[i].get_text()
            extracted_pages.append((i + 1, text))
    elif ext == ".docx":
        docx_doc = docx.Document(io.BytesIO(content))
        full_text = "\n".join([p.text for p in docx_doc.paragraphs if p.text.strip()])
        extracted_pages.append((1, full_text))
    elif ext in [".txt", ".md"]:
        text = content.decode("utf-8", errors="replace")
        extracted_pages.append((1, text))
    else:
        raise ValueError(f"Unsupported file format: {ext}")
        
    return extracted_pages

def chunk_text(pages: List[Tuple[int, str]], chunk_size_words: int = 350, overlap_words: int = 35) -> List[Dict[str, Any]]:
    """Chunks text while preserving page number references."""
    chunks = []
    chunk_index = 1
    
    for page_num, page_text in pages:
        words = page_text.split()
        if not words:
            continue
        
        i = 0
        while i < len(words):
            chunk_words = words[i:i + chunk_size_words]
            chunk_str = " ".join(chunk_words)
            if chunk_str.strip():
                chunks.append({
                    "chunk_seq": chunk_index,
                    "page": page_num,
                    "text": chunk_str
                })
                chunk_index += 1
            if i + chunk_size_words >= len(words):
                break
            i += (chunk_size_words - overlap_words)
            
    return chunks

async def ingest_document(
    filename: str,
    content: bytes,
    title: str,
    department: str,
    classification: str,
    tenant_id: str,
    user_id: int,
    vendor: str = "",
    tags: List[str] = None,
    custom_fields: Dict[str, Any] = None,
    access_policies: List[Dict[str, Any]] = None,
    db = None
) -> Tuple[Document, Dict[str, Any]]:
    doc_id = f"DOC-{uuid.uuid4().hex[:8].upper()}"
    tags = tags or []
    custom_fields = custom_fields or {}
    access_policies = access_policies or []

    # 1. Validation
    if not content or len(content) == 0:
        raise ValueError("Uploaded file is empty")
    
    ext = os.path.splitext(filename)[1].lower()
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise ValueError(f"Invalid file type {ext}. Allowed: {settings.ALLOWED_EXTENSIONS}")

    # 2. Extract Text
    pages = extract_text_from_file(filename, content)
    total_text = "".join([p[1] for p in pages]).strip()
    if not total_text:
        raise ValueError("No extractable text found in document")

    # 3. Store MongoDB Metadata
    mongo_db = get_mongo_db()
    mongo_record = {
        "document_id": doc_id,
        "title": title,
        "department": department.lower(),
        "classification": classification.lower(),
        "vendor": vendor,
        "tags": tags,
        "custom_fields": custom_fields,
        "created_at": datetime.datetime.now(datetime.timezone.utc),
        "page_count": len(pages),
        "raw_character_count": len(total_text)
    }
    if mongo_db is not None:
        await mongo_db.documents.insert_one(mongo_record)
        await save_mongo_mock_backup()

    # 4. Create PostgreSQL Record
    dept_id = None
    if db is not None:
        from sqlalchemy import select
        from app.models.postgres_models import Department
        dept_res = await db.execute(select(Department).where(Department.name == department.lower()))
        dept_obj = dept_res.scalars().first()
        if dept_obj:
            dept_id = dept_obj.id

    pg_doc = Document(
        id=doc_id,
        mongo_document_id=doc_id,
        title=title,
        status="DRAFT",
        tenant_id=tenant_id,
        department_id=dept_id,
        department=department.lower(),
        classification=classification.lower(),
        created_by=user_id
    )
    db.add(pg_doc)
    await db.commit()
    await db.refresh(pg_doc)

    # 4b. Explicit Document Access Policies (ACLs)
    if access_policies:
        for ap in access_policies:
            policy_rec = DocumentAccessPolicy(
                document_id=doc_id,
                grantee_type=ap.get("grantee_type", "USER").upper(),
                grantee_id=str(ap.get("grantee_id", "")),
                permission_level=ap.get("permission_level", "READ").upper()
            )
            db.add(policy_rec)
        await db.commit()

    # 5. Chunking & PostgreSQL pgvector Ingestion
    chunks = chunk_text(pages)
    if chunks:
        model = get_embedding_model()
        chunk_texts = [c["text"] for c in chunks]
        embeddings = model.encode(chunk_texts, show_progress_bar=False)

        for idx, c in enumerate(chunks):
            c["chunk_id"] = f"{doc_id}-P{c['page']}-{c['chunk_seq']}"
            c["embedding"] = embeddings[idx].tolist()

            pg_chunk = DocumentChunk(
                id=c["chunk_id"],
                document_id=doc_id,
                chunk_seq=c["chunk_seq"],
                page=c["page"],
                text=c["text"],
                embedding=c["embedding"],
                tsv=func.to_tsvector("english", c["text"])
            )
            db.add(pg_chunk)
        await db.commit()

        # Optional legacy ChromaDB sync
        try:
            col = get_chroma_collection("document_chunks")
            if col:
                col.add(
                    ids=[c["chunk_id"] for c in chunks],
                    documents=[c["text"] for c in chunks],
                    embeddings=[c["embedding"] for c in chunks],
                    metadatas=[{
                        "document_id": doc_id,
                        "document_title": title,
                        "tenant_id": tenant_id,
                        "department": department.lower(),
                        "classification": classification.lower(),
                        "page": int(c["page"]),
                        "chunk_id": c["chunk_id"]
                    } for c in chunks]
                )
        except Exception as e:
            logger.warning(f"ChromaDB legacy sync notice: {e}")

    # 7. Audit Events
    await log_audit_event(
        db,
        action="DOCUMENT_UPLOADED",
        actor_id=user_id,
        document_id=doc_id,
        metadata_info={"filename": filename, "size_bytes": len(content)}
    )
    await log_audit_event(
        db,
        action="DOCUMENT_INGESTED",
        actor_id=user_id,
        document_id=doc_id,
        metadata_info={"chunks_count": len(chunks), "pages": len(pages), "department": department}
    )

    # 8. Publish Kafka / Event Bus Event
    try:
        from app.events.kafka_producer import publish_workflow_event
        await publish_workflow_event(
            event_type="DOCUMENT_UPLOADED",
            document_id=doc_id,
            document_title=title,
            department=department.lower(),
            from_status=None,
            to_status="DRAFT",
            actor_id=user_id,
            tenant_id=tenant_id
        )
    except Exception as e:
        logger.warning(f"Failed to publish document upload event to Kafka: {e}")

    return pg_doc, mongo_record
